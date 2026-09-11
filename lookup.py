"""
Look up rota data by hospital + date, straight out of the Markdown transcripts.

get_hospital_date_rows(hospital, date, transcripts_dir) searches every .md
transcript for pages that (a) mention `hospital` anywhere in the page and
(b) are stated to cover `date`'s month/year, then — within those pages —
finds the table column whose header matches `date`'s day-of-month and pulls
out every row's value in that column.

hospital_date_column(...) builds on that for the common case (one clearly
matching page): an ordered, human-readable list of {label, value} for that
date's column, the RATIO row's value pulled out separately, and warnings on
any row whose data looks shifted out of place (see _flag_misaligned below).

export_column_xlsx(...) writes that same column out as a small styled .xlsx.

Rows come back exactly as transcribed (raw values, uninterpreted) since this
reads the transcripts directly rather than the cleaned-up .xlsx output —
deliberately "give me everything tied to this date" rather than a processed
answer.

get_hospital_date_rows_from_xlsx(...) / hospital_date_column_from_xlsx(...)
do the same lookup, but against an already-built workbook (one produced by
to_xlsx.convert_all()) instead of the raw .md files. Use these once the
Markdown transcripts have been converted: since to_xlsx.py already ran every
table through the same format-tolerant parsing (merged split headers and
RATIO rows, [red]/[blank]/<br>/** noise turned into real Excel formatting or
dropped), every sheet in the workbook is in one consistent shape no matter
which prompt/model-format produced the original page — so looking things up
there is simpler and a bit cleaner than re-parsing Markdown each time.

Why search-by-substring instead of a hospital ID: nothing in this project
currently assigns hospitals a code — config.FACILITY is a plain string
("Kakamega NBU"), and each page's own header text says which hospital it's
from (e.g. "COUNTY GOVERNMENT OF KAKAMEGA"). Matching that text directly
means this works today, across however many hospitals' transcripts end up
sharing a transcripts folder, without inventing an ID scheme first.

This version favors plain for-loops and if/else statements over shorter
one-liners (list comprehensions, ternary expressions) even where a
one-liner would do the same thing, to keep each step easy to follow.
"""
from __future__ import annotations

import re
from datetime import date as date_cls
from pathlib import Path
from typing import Union

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.comments import Comment

from to_xlsx import split_segments, parse_table_rows, build_header_and_data  # reuse the existing parser

MONTH_YEAR_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)"
    r"\s+(\d{4})",
    re.IGNORECASE,
)
MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}
LEADING_DAY_RE = re.compile(r"^\D*(\d{1,2})\b")

# How many letters "off" a hospital-name match is still allowed to be, e.g.
# matching the typo "KAKAMEA" (one letter dropped) against "Kakamega".
HOSPITAL_NAME_MAX_TYPOS = 2

# Recognized shift codes, per the transcription prompt's conventions
# (M/N/No/D/Do/PH/E) plus the special-row designations seen in practice
# (I/C, DEP I/C). Used only to flag rows that look shifted out of place —
# see _flag_misaligned — never to "correct" a value.
KNOWN_SHIFT_CODES = {"M", "E", "N", "D", "DO", "NO", "PH", "I/C", "DEP I/C"}

FONT_NAME = "Arial"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _as_date(value: Union[str, date_cls]) -> date_cls:
    """Accept either a real date or an ISO string like '2022-01-10'."""
    if isinstance(value, date_cls):
        return value
    return date_cls.fromisoformat(str(value))


def _edit_distance(a: str, b: str) -> int:
    """Plain Levenshtein edit distance between two strings (single-character
    inserts, deletes, and substitutions each count as 1). Used only to
    tolerate small OCR/transcription typos in a hospital name — never to
    "correct" any rota data itself.
    """
    rows = len(a) + 1
    cols = len(b) + 1
    table = []
    for i in range(rows):
        table.append([0] * cols)
    for i in range(rows):
        table[i][0] = i
    for j in range(cols):
        table[0][j] = j

    for i in range(1, rows):
        for j in range(1, cols):
            if a[i - 1] == b[j - 1]:
                cost = 0
            else:
                cost = 1
            delete_cost = table[i - 1][j] + 1
            insert_cost = table[i][j - 1] + 1
            substitute_cost = table[i - 1][j - 1] + cost
            table[i][j] = min(delete_cost, insert_cost, substitute_cost)

    return table[rows - 1][cols - 1]


def _hospital_mentioned(hospital: str, md_text: str) -> bool:
    """True if `hospital` appears in `md_text`, allowing for a small typo.

    Tries an exact case-insensitive substring match first (the common,
    cheap case). If that fails, slides a same-ish-length window across the
    page text and checks whether any window is within
    HOSPITAL_NAME_MAX_TYPOS single-character edits of `hospital` — this is
    what lets "Kakamega" match a page that transcribed it as "KAKAMEA"
    (one letter dropped).
    """
    hospital_lower = hospital.lower()
    text_lower = md_text.lower()

    if hospital_lower in text_lower:
        return True

    needle_len = len(hospital_lower)
    shortest_window = needle_len - HOSPITAL_NAME_MAX_TYPOS
    longest_window = needle_len + HOSPITAL_NAME_MAX_TYPOS
    if shortest_window < 1:
        shortest_window = 1

    window_len = shortest_window
    while window_len <= longest_window:
        start = 0
        while start + window_len <= len(text_lower):
            window = text_lower[start:start + window_len]
            if _edit_distance(window, hospital_lower) <= HOSPITAL_NAME_MAX_TYPOS:
                return True
            start += 1
        window_len += 1

    return False


def _page_month_year(md_text: str) -> tuple[int, int] | None:
    """Pull (month, year) from anywhere on the page, e.g. '...ROTA FOR JANUARY 2022'.

    Returns None if the page doesn't state a month/year anywhere — in that
    case the caller treats it as "can't rule this page out" and checks its
    date columns anyway, rather than silently skipping it.
    """
    m = MONTH_YEAR_RE.search(md_text)
    if not m:
        return None
    return MONTHS[m.group(1).lower()], int(m.group(2))


def _matching_day_column(header: list[str], day: int) -> int | None:
    """Index of the header cell whose leading number equals `day` (e.g. '10 MON' -> 10)."""
    for j, cell in enumerate(header):
        m = LEADING_DAY_RE.match(cell.strip())
        if m and int(m.group(1)) == day:
            return j
    return None


def _row_label(header: list[str], row: list[str]) -> str:
    """Pick the most useful identifier for a row: the DATE-DAY-style label
    column if present (e.g. "DEP I/C", "RATIO"), else the NAMES column,
    else "NO <n>". Real transcripts here usually have NAMES blank (the
    scanned form has the names column physically redacted), so DATE-DAY /
    NO end up doing most of the identifying work in practice.
    """
    if len(row) > 0:
        no = row[0].strip()
    else:
        no = ""

    if len(row) > 1:
        names = row[1].strip()
    else:
        names = ""

    if len(row) > 2:
        date_day = row[2].strip()
    else:
        date_day = ""

    if date_day:
        return date_day
    if names:
        return names
    if no:
        return f"NO {no}"
    return "?"


def _is_ratio_row(row: list[str]) -> bool:
    """True for the RATIO/summary row. Checks the first few columns for
    either "ratio" or "support staff" (not just "ratio" alone), since some
    transcripts' RATIO row ends up with "SUPPORT STAFF" in the column that
    other transcripts use for "RATIO" itself once a split RATIO row has
    been merged back together (see merge_split_data_rows).
    """
    for c in row[:3]:
        text = c.strip().lower()
        if "ratio" in text or "support staff" in text:
            return True
    return False


def _phone_column_index(header: list[str]) -> int | None:
    """Find the PHONE NO column by name rather than assuming it's the last
    one — parse_table_rows pads every row to the widest row in the table,
    which can leave one or more truly-blank trailing columns after it.
    """
    for i, cell in enumerate(header):
        if "phone" in cell.lower():
            return i
    return None


def _flag_misaligned(header: list[str], row: list[str]) -> bool:
    """Heuristic: the PHONE NO column, in every real transcript seen so far,
    is either blank or an actual phone number. If it instead holds
    something that reads like a shift code (M/E/N/Do/No/PH/I/C/...), that's
    a strong sign the model shifted this row's cells one column to the
    right somewhere earlier — seen in practice on rows re-numbering back to
    "1"/"2" at the bottom of a page, where an extra stray cell gets
    inserted. This flags it rather than guessing a fix, since silently
    "correcting" a value risks being wrong in a way that's much harder to
    notice than a raw misalignment is.
    """
    phone_idx = _phone_column_index(header)
    if phone_idx is None or phone_idx >= len(row):
        return False
    val = row[phone_idx].strip()
    return bool(val) and val.upper() in KNOWN_SHIFT_CODES


def get_hospital_date_rows(
    hospital: str,
    date: Union[str, date_cls],
    transcripts_dir: Path,
) -> list[dict]:
    """Return every rota row tied to `hospital` + `date` across all transcripts.

    hospital: case-insensitive substring matched against each transcript's
              full text (e.g. "Kakamega" matches "COUNTY GOVERNMENT OF
              KAKAMEGA"). Also tolerates a small typo either side (up to
              HOSPITAL_NAME_MAX_TYPOS single-character edits), since some
              transcripts render this as "KAKAMEA" (a dropped letter) —
              see _hospital_mentioned.
    date:     a datetime.date, or an ISO string like "2022-01-10".
    transcripts_dir: folder containing the .md transcripts to search
                      (typically config.TRANSCRIPTS_DIR).

    Returns a list of dicts, one per matching row, each with:
        source_file  - which .md file this came from
        table_index  - which table on that page (0-based, in page order)
        header       - the full column header row (already merged if the
                       page split it across two physical rows)
        row          - the full row, as transcribed (raw, unmodified values)
        date_column  - the header text of the matched date column
        value        - just that row's value in the matched date column
        label        - a human-readable row identifier (see _row_label)
        is_ratio     - True for the RATIO/summary row, not an individual staff row
        misaligned   - True if this row's cells look shifted (see _flag_misaligned) —
                       `value` may not be trustworthy for a flagged row

    A page whose month/year (parsed from its own text) doesn't match
    `date` is skipped outright, so e.g. a "10" day column on a February
    page never gets confused with January 10th. Returns [] if nothing
    matches — check that transcripts_dir actually holds the pages you
    expect before assuming a real absence.
    """
    target = _as_date(date)
    transcripts_dir = Path(transcripts_dir)
    results: list[dict] = []

    for md_path in sorted(transcripts_dir.glob("*.md")):
        md_text = md_path.read_text(encoding="utf-8")

        if not _hospital_mentioned(hospital, md_text):
            continue

        page_month_year = _page_month_year(md_text)
        if page_month_year is not None and page_month_year != (target.month, target.year):
            continue  # this page is a different month/year — its day numbers don't apply

        table_index = -1
        for kind, lines in split_segments(md_text):
            if kind != "table":
                continue
            table_index += 1

            rows = parse_table_rows(lines)
            if not rows:
                continue
            header, data_rows = build_header_and_data(rows)

            col = _matching_day_column(header, target.day)
            if col is None:
                continue

            for row in data_rows:
                has_any_content = False
                for c in row:
                    if c.strip():
                        has_any_content = True
                        break
                if not has_any_content:
                    continue  # a fully blank filler row — not real data

                if col >= len(row):
                    continue

                results.append({
                    "source_file": md_path.name,
                    "table_index": table_index,
                    "header": header,
                    "row": row,
                    "date_column": header[col],
                    "value": row[col],
                    "label": _row_label(header, row),
                    "is_ratio": _is_ratio_row(row),
                    "misaligned": _flag_misaligned(header, row),
                })

    return results


def hospital_date_column(
    hospital: str,
    date: Union[str, date_cls],
    transcripts_dir: Path,
    prompt_name: str | None = None,
) -> dict:
    """Convenience wrapper for the common case: one page clearly matches.

    prompt_name: optional filter for when main.py saved more than one
        prompt's output for the same page (e.g. page_001_base.md and
        page_001_base2.md, per storage.py's "<page>_<prompt_name>.md"
        naming). Pass just the prompt part, e.g. "base2", and only files
        whose name ends in "_base2.md" are considered. Leave it out if you
        only ever ran one prompt against this page.

    Returns:
        {
          "source_file": "page_003.md",
          "date_column": "10 MON",
          "entries": [{"label": "WC", "value": "S", "misaligned": False}, ...],
          "ratio": "2/2/2" or None,
          "warnings": ["...one line per flagged row..."],
        }

    Raises ValueError if zero, or more than one, source file matches after
    the prompt_name filter (if given) — check get_hospital_date_rows
    directly if you want to see every matching file instead of picking one.
    """
    rows = get_hospital_date_rows(hospital, date, transcripts_dir)

    if prompt_name:
        filtered_rows = []
        for r in rows:
            file_stem = Path(r["source_file"]).stem  # e.g. "page_001_base2"
            if file_stem.endswith(f"_{prompt_name}"):
                filtered_rows.append(r)
        rows = filtered_rows

    if not rows:
        if prompt_name:
            raise ValueError(
                f"No data found for hospital={hospital!r}, date={date!r}, "
                f"prompt_name={prompt_name!r} in {transcripts_dir}"
            )
        raise ValueError(f"No data found for hospital={hospital!r}, date={date!r} in {transcripts_dir}")

    files = set()
    for r in rows:
        files.add(r["source_file"])

    if len(files) > 1:
        raise ValueError(
            f"{len(files)} different transcript files matched ({sorted(files)}) — "
            f"pass prompt_name= to pick one, e.g. prompt_name=\"base2\" for page_001_base2.md."
        )

    entries = []
    ratio = None
    warnings = []
    for r in rows:
        if r["is_ratio"]:
            ratio = r["value"]
            continue
        entries.append({"label": r["label"], "value": r["value"], "misaligned": r["misaligned"]})
        if r["misaligned"]:
            warnings.append(
                f"Row {r['label']!r}: value {r['value']!r} looks shifted out of place "
                f"(its PHONE NO column held a shift-code-like value) — verify against the source page."
            )

    return {
        "source_file": rows[0]["source_file"],
        "date_column": rows[0]["date_column"],
        "entries": entries,
        "ratio": ratio,
        "warnings": warnings,
    }


def export_column_xlsx(
    hospital: str,
    date: Union[str, date_cls],
    transcripts_dir: Path,
    output_path: Path,
    prompt_name: str | None = None,
) -> Path:
    """Write hospital_date_column(...)'s result out as a small styled .xlsx.

    One row per entry (Label, Value), the RATIO value as its own bold row at
    the bottom, and a cell comment on any row flagged as possibly misaligned.
    prompt_name is passed straight through to hospital_date_column — see its
    docstring if more than one prompt's output exists for the same page.
    """
    result = hospital_date_column(hospital, date, transcripts_dir, prompt_name=prompt_name)

    wb = openpyxl.Workbook()
    ws = wb.active

    raw_title = f"{hospital}_{date}"
    cleaned_title = re.sub(r"[\[\]:*?/\\]", "_", raw_title)
    ws.title = cleaned_title[:31]  # Excel sheet names can't be longer than 31 characters

    ws.merge_cells("A1:B1")
    title_cell = ws.cell(row=1, column=1, value=f"{hospital} — {result['date_column']} ({date})")
    title_cell.font = Font(name=FONT_NAME, bold=True, size=12)
    title_cell.alignment = CENTER

    header_row = 3
    for j, text in enumerate(("Label", "Value"), start=1):
        cell = ws.cell(row=header_row, column=j, value=text)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.alignment = CENTER
        cell.border = BORDER

    r = header_row + 1
    for entry in result["entries"]:
        label_cell = ws.cell(row=r, column=1, value=entry["label"])
        value_cell = ws.cell(row=r, column=2, value=entry["value"])
        for cell in (label_cell, value_cell):
            cell.font = Font(name=FONT_NAME)
            cell.alignment = CENTER
            cell.border = BORDER
        if entry["misaligned"]:
            value_cell.comment = Comment(
                "Possibly misaligned in the source transcript — verify against the scanned page.",
                "Rota Lookup",
            )
        r += 1

    if result["ratio"] is not None:
        r += 1
        label_cell = ws.cell(row=r, column=1, value="RATIO")
        value_cell = ws.cell(row=r, column=2, value=result["ratio"])
        for cell in (label_cell, value_cell):
            cell.font = Font(name=FONT_NAME, bold=True)
            cell.alignment = CENTER
            cell.border = BORDER

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Same lookup, but reading an already-built .xlsx instead of raw Markdown
# ---------------------------------------------------------------------------

def _sheet_text(ws) -> str:
    """Collect every non-empty cell's text on a worksheet into one string,
    for the same "does this page mention this hospital / month" checks used
    against raw Markdown text.
    """
    pieces = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                pieces.append(str(cell.value))
    return "\n".join(pieces)


def _cell_has_table_border(cell) -> bool:
    """True if a cell carries the thin border to_xlsx.render_table() puts on
    every table cell (see to_xlsx.BORDER) — used to tell a table row apart
    from an ordinary heading/paragraph row, see _find_table_blocks_in_sheet
    for why this is checked instead of the cell's value.
    """
    border = cell.border
    if border is None:
        return False
    if border.left is not None and border.left.style is not None:
        return True
    if border.top is not None and border.top.style is not None:
        return True
    return False


def _find_table_blocks_in_sheet(ws) -> list[list[list[str]]]:
    """Find each table's rows (header + data, already merged) inside one
    sheet of a workbook built by to_xlsx.convert_all().

    A row counts as "part of a table" if any of its cells carries the thin
    border to_xlsx.render_table() applies to every table cell, header and
    data alike — including a cell whose value is an intentional blank.
    That border is what's checked here rather than the cell's value: a cell
    explicitly set to "" round-trips through a saved-and-reopened workbook
    as None, exactly like a column to_xlsx.render_text_block() never wrote
    to at all, so value alone can't tell a mostly-blank table row apart
    from an ordinary heading/paragraph row — but the border survives that
    round trip.

    Returns a list of blocks, one per table on the sheet, each block being
    a list of rows (a list of cell-text strings) in the order they appear —
    block[0] is that table's header row, block[1:] are its data rows.
    """
    max_row = ws.max_row
    max_col = ws.max_column

    row_texts = []
    row_is_table = []
    for r in range(1, max_row + 1):
        texts = []
        any_bordered = False
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is None:
                texts.append("")
            else:
                texts.append(str(cell.value))
            if _cell_has_table_border(cell):
                any_bordered = True
        row_texts.append(texts)
        row_is_table.append(any_bordered)

    blocks = []
    current_block = []
    for i in range(max_row):
        if row_is_table[i]:
            current_block.append(row_texts[i])
        else:
            if current_block:
                blocks.append(current_block)
                current_block = []
    if current_block:
        blocks.append(current_block)
    return blocks


def get_hospital_date_rows_from_xlsx(
    hospital: str,
    date: Union[str, date_cls],
    xlsx_path: Path,
) -> list[dict]:
    """Same idea as get_hospital_date_rows, but reads an already-built
    workbook (one produced by to_xlsx.convert_all()) instead of re-parsing
    the raw Markdown transcripts. See the module docstring for why this
    version can be simpler and cleaner: the format-tolerant parsing already
    happened once, when the workbook was built.

    Returns the same list-of-dicts shape as get_hospital_date_rows, except:
      - "source_file" holds the sheet name (e.g. "page_001_base2") instead
        of a .md filename.
      - "misaligned" is always False. The PHONE-NO-column heuristic that
        flags a shifted row needs the row's original raw text, and by the
        time a value is sitting in a cell here, [red]/[blank]/other tags
        have already been resolved into formatting or dropped — so that
        check is only meaningful against the raw Markdown, not the xlsx.
    """
    target = _as_date(date)
    xlsx_path = Path(xlsx_path)
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    results: list[dict] = []

    for ws in workbook.worksheets:
        sheet_text = _sheet_text(ws)

        if not _hospital_mentioned(hospital, sheet_text):
            continue

        page_month_year = _page_month_year(sheet_text)
        if page_month_year is not None and page_month_year != (target.month, target.year):
            continue  # this sheet is a different month/year — its day numbers don't apply

        blocks = _find_table_blocks_in_sheet(ws)
        for table_index, block in enumerate(blocks):
            if len(block) < 2:
                continue  # a table needs at least a header row and one data row

            header = block[0]
            data_rows = block[1:]

            col = _matching_day_column(header, target.day)
            if col is None:
                continue

            for row in data_rows:
                has_any_content = False
                for c in row:
                    if c.strip():
                        has_any_content = True
                        break
                if not has_any_content:
                    continue  # a fully blank filler row — not real data

                if col >= len(row):
                    continue

                results.append({
                    "source_file": ws.title,
                    "table_index": table_index,
                    "header": header,
                    "row": row,
                    "date_column": header[col],
                    "value": row[col],
                    "label": _row_label(header, row),
                    "is_ratio": _is_ratio_row(row),
                    "misaligned": False,
                })

    return results


def hospital_date_column_from_xlsx(
    hospital: str,
    date: Union[str, date_cls],
    xlsx_path: Path,
    sheet_name: str | None = None,
) -> list[dict]:
    """Convenience wrapper around get_hospital_date_rows_from_xlsx: pulls out
    one {entries, ratio} result per sheet that matches hospital + date.

    Unlike hospital_date_column() (the Markdown version), this does NOT
    require exactly one file to match, so it never raises just because more
    than one sheet has data for that hospital/date. Real workbooks often
    carry more than one prompt's output for the same page while you're
    still comparing prompts (e.g. "page_046" and "page_046A") — this just
    returns a result for each one, in sheet-name order, so "page_046" comes
    before "page_046A". Once you delete the extra sheet/file, the exact
    same call naturally returns a list of just one result — nothing about
    how you call this needs to change either way.

    sheet_name: optional filter to see just one specific sheet instead of
        every match (e.g. "page_001_base2").

    Returns a list, one item per matching sheet:
        [
          {
            "source_file": "page_046",
            "date_column": "9 SUN",
            "entries": [{"label": "I/C", "value": "SI", "misaligned": False}, ...],
            "ratio": "3/1/2" or None,
            "warnings": [],   # always empty here — see get_hospital_date_rows_from_xlsx
          },
          {
            "source_file": "page_046A",
            ...
          },
        ]

    Raises ValueError only if nothing at all matches (including after the
    sheet_name filter, if given).
    """
    rows = get_hospital_date_rows_from_xlsx(hospital, date, xlsx_path)

    if sheet_name:
        filtered_rows = []
        for r in rows:
            if r["source_file"] == sheet_name:
                filtered_rows.append(r)
        rows = filtered_rows

    if not rows:
        if sheet_name:
            raise ValueError(
                f"No data found for hospital={hospital!r}, date={date!r}, "
                f"sheet_name={sheet_name!r} in {xlsx_path}"
            )
        raise ValueError(f"No data found for hospital={hospital!r}, date={date!r} in {xlsx_path}")

    sheets_in_order = []
    for r in rows:
        if r["source_file"] not in sheets_in_order:
            sheets_in_order.append(r["source_file"])
    sheets_in_order.sort()

    results = []
    for sheet in sheets_in_order:
        sheet_rows = []
        for r in rows:
            if r["source_file"] == sheet:
                sheet_rows.append(r)

        entries = []
        ratio = None
        for r in sheet_rows:
            if r["is_ratio"]:
                ratio = r["value"]
                continue
            entries.append({"label": r["label"], "value": r["value"], "misaligned": r["misaligned"]})

        results.append({
            "source_file": sheet,
            "date_column": sheet_rows[0]["date_column"],
            "entries": entries,
            "ratio": ratio,
            "warnings": [],
        })

    return results


if __name__ == "__main__":
    import sys
    import config

    if len(sys.argv) not in (3, 4):
        raise SystemExit(
            "Usage: python lookup.py <hospital substring> <YYYY-MM-DD> [prompt_name or xlsx_path]\n"
            "  The third argument is optional.\n"
            "  - If it ends in \".xlsx\" (e.g. an output of to_xlsx.py), that workbook\n"
            "    is looked up directly instead of the raw Markdown transcripts.\n"
            "  - Otherwise it's treated as a prompt_name filter — only needed if more\n"
            "    than one prompt's output exists for the same page, e.g. \"base2\" for\n"
            "    page_001_base2.md."
        )

    hospital_arg, date_arg = sys.argv[1], sys.argv[2]
    if len(sys.argv) == 4:
        third_arg = sys.argv[3]
    else:
        third_arg = None

    def print_one_result(result):
        print(f"{hospital_arg} — {result['date_column']} ({date_arg})  [{result['source_file']}]")
        for entry in result["entries"]:
            if entry["misaligned"]:
                flag = "  [!]"
            else:
                flag = ""
            print(f"  {entry['label']:<10} {entry['value']!r}{flag}")
        print(f"  {'RATIO':<10} {result['ratio']!r}")
        for w in result["warnings"]:
            print(f"[WARNING] {w}")

    if third_arg and third_arg.lower().endswith(".xlsx"):
        try:
            # one result per matching sheet — usually just one, but can be
            # more while you're still comparing more than one prompt's
            # output for the same page (see hospital_date_column_from_xlsx)
            results = hospital_date_column_from_xlsx(hospital_arg, date_arg, third_arg)
        except ValueError as e:
            raise SystemExit(str(e))

        first = True
        for result in results:
            if not first:
                print()
            first = False
            print(f"=== {result['source_file']} ===")
            print_one_result(result)
    else:
        try:
            result = hospital_date_column(hospital_arg, date_arg, config.TRANSCRIPTS_DIR, prompt_name=third_arg)
        except ValueError as e:
            raise SystemExit(str(e))

        print_one_result(result)